"""Live Excel report: one row per page, saved after every page."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

COLUMNS = [
    "Folder Name",
    "File Name",
    "Page",
    "Input DPI",
    "DPI Source",
    "Width px",
    "Height px",
    "Quality Score",
    "Quality",
    "Quality Warning",
    "Sharpness",
    "Blur",
    "Contrast",
    "Noise",
    "Brightness",
    "Clipping",
    "Text Visibility",
    "Compression",
    "Usable Dimension",
    "Effective Resolution",
    "Document Type",
    "P(Handwritten)",
    "Document Type Method",
    "Final Status",
    "Error",
    "Time taken (s)",
]

_STATUS_FILL = {
    "OK": PatternFill("solid", fgColor="C6EFCE"),
    "REVIEW_REQUIRED": PatternFill("solid", fgColor="FFEB9C"),
    "ERROR": PatternFill("solid", fgColor="FFC7CE"),
}


class ExcelReport:
    def __init__(self, path: Path, *, resume: bool = False) -> None:
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if resume and self.path.is_file():
            self.wb = load_workbook(self.path)
            self.ws = self.wb.active
        else:
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = "Pages"
            self.ws.append(COLUMNS)
            for cell in self.ws[1]:
                cell.font = Font(bold=True)
            self.ws.freeze_panes = "A2"
            self.save()
        self._status_col = COLUMNS.index("Final Status") + 1

    def add(self, row: dict[str, Any]) -> None:
        self.ws.append([row.get(col) for col in COLUMNS])
        fill = _STATUS_FILL.get(str(row.get("Final Status")))
        if fill is not None:
            self.ws.cell(row=self.ws.max_row, column=self._status_col).fill = fill
        self.save()

    def save(self) -> None:
        try:
            self.wb.save(self.path)
        except PermissionError:
            # Excel holds a lock while the file is open; the next page retries.
            pass

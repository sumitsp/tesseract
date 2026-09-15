"""Professional Excel report: one row per page plus a Summary sheet."""

from __future__ import annotations

from pathlib import Path
from statistics import mean

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from image_preprocessing.results import PageResult

COLUMNS = [
    ("document_name", 28),
    ("page_number", 14),
    ("input_file", 42),
    ("input_format", 14),
    ("width_px", 12),
    ("height_px", 12),
    ("input_dpi", 12),
    ("output_dpi", 12),
    ("dpi_action", 26),
    ("dpi_source", 22),
    ("quality_score", 14),
    ("quality_warning", 28),
    ("sharpness_score", 16),
    ("contrast_score", 16),
    ("noise_score", 14),
    ("document_type", 16),
    ("document_type_confidence", 24),
    ("document_type_method", 20),
    ("rotation_angle", 16),
    ("rotation_confidence", 20),
    ("rotation_status", 18),
    ("mirror", 12),
    ("mirror_confidence", 18),
    ("mirror_corrected", 18),
    ("tilt_angle", 12),
    ("tilt_confidence", 16),
    ("tilt_status", 14),
    ("final_status", 18),
    ("warnings", 48),
    ("output_file", 42),
    ("processing_time_seconds", 22),
    ("error_message", 32),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
ALT_FILL = PatternFill("solid", fgColor="F2F2F2")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
REVIEW_FILL = PatternFill("solid", fgColor="F8CBAD")
ERROR_FILL = PatternFill("solid", fgColor="F4C7C3")
OK_FILL = PatternFill("solid", fgColor="C6EFCE")
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
ANGLE_FORMAT = "0.00"
CONF_FORMAT = "0.0000"
DPI_FORMAT = "0.0"
INT_FORMAT = "0"
TIME_FORMAT = "0.000"


def _value(result: PageResult, key: str):
    return getattr(result, key, None)


def write_excel_report(path: Path, results: list[PageResult], quality_threshold: float) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    pages = wb.active
    pages.title = "Pages"
    _write_pages_sheet(pages, results)
    summary = wb.create_sheet("Summary")
    _write_summary_sheet(summary, results, quality_threshold)
    wb.save(path)
    return path


def _write_pages_sheet(ws: Worksheet, results: list[PageResult]) -> None:
    keys = [k for k, _ in COLUMNS]
    for col, (key, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(1, col, key)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(2, len(results) + 1)}"
    ws.row_dimensions[1].height = 22

    angle_cols = {"rotation_angle", "tilt_angle"}
    conf_cols = {
        "rotation_confidence",
        "mirror_confidence",
        "tilt_confidence",
        "document_type_confidence",
    }
    dpi_cols = {"input_dpi", "output_dpi", "quality_score", "sharpness_score", "contrast_score", "noise_score"}

    for row_i, result in enumerate(results, start=2):
        for col, key in enumerate(keys, start=1):
            value = _value(result, key)
            cell = ws.cell(row_i, col, value)
            cell.border = THIN
            cell.alignment = Alignment(vertical="center", wrap_text=key in {"warnings", "input_file", "output_file"})
            if row_i % 2 == 0:
                cell.fill = ALT_FILL
            if key in angle_cols and isinstance(value, (int, float)):
                cell.number_format = ANGLE_FORMAT
            elif key in conf_cols and isinstance(value, (int, float)):
                cell.number_format = CONF_FORMAT
            elif key in dpi_cols and isinstance(value, (int, float)):
                cell.number_format = DPI_FORMAT
            elif key in {"width_px", "height_px", "page_number"} and isinstance(value, (int, float)):
                cell.number_format = INT_FORMAT
            elif key == "processing_time_seconds" and isinstance(value, (int, float)):
                cell.number_format = TIME_FORMAT

    last_row = max(2, len(results) + 1)
    last_col = get_column_letter(len(keys))
    status_col = get_column_letter(keys.index("final_status") + 1)
    warn_col = get_column_letter(keys.index("warnings") + 1)
    rot_status = get_column_letter(keys.index("rotation_status") + 1)
    tilt_status = get_column_letter(keys.index("tilt_status") + 1)
    mirror_col = get_column_letter(keys.index("mirror") + 1)

    ws.conditional_formatting.add(
        f"{status_col}2:{status_col}{last_row}",
        CellIsRule(operator="equal", formula=['"REVIEW_REQUIRED"'], fill=REVIEW_FILL),
    )
    ws.conditional_formatting.add(
        f"{status_col}2:{status_col}{last_row}",
        CellIsRule(operator="equal", formula=['"ERROR"'], fill=ERROR_FILL),
    )
    ws.conditional_formatting.add(
        f"{status_col}2:{status_col}{last_row}",
        CellIsRule(operator="equal", formula=['"CORRECTED"'], fill=OK_FILL),
    )
    ws.conditional_formatting.add(
        f"A2:{last_col}{last_row}",
        FormulaRule(formula=[f'LEN({warn_col}2)>0'], fill=WARN_FILL),
    )
    ws.conditional_formatting.add(
        f"{rot_status}2:{rot_status}{last_row}",
        CellIsRule(operator="equal", formula=['"UNCERTAIN"'], fill=WARN_FILL),
    )
    ws.conditional_formatting.add(
        f"{tilt_status}2:{tilt_status}{last_row}",
        CellIsRule(operator="equal", formula=['"UNCERTAIN"'], fill=WARN_FILL),
    )
    ws.conditional_formatting.add(
        f"{mirror_col}2:{mirror_col}{last_row}",
        CellIsRule(operator="equal", formula=['"UNKNOWN"'], fill=WARN_FILL),
    )


def _count(results: list[PageResult], pred) -> int:
    return sum(1 for r in results if pred(r))


def _write_summary_sheet(ws: Worksheet, results: list[PageResult], quality_threshold: float) -> None:
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18
    title = ws.cell(1, 1, "Preprocessing summary")
    title.font = Font(bold=True, size=14, color="1F4E79", name="Calibri")
    ws.merge_cells("A1:B1")

    docs = {r.document_name for r in results}
    qualities = [r.quality_score for r in results if r.quality_score is not None]
    rows = [
        ("Total documents", len(docs)),
        ("Total pages", len(results)),
        (
            "Pages successfully processed",
            _count(results, lambda r: r.final_status in {"CORRECTED", "UNCHANGED"}),
        ),
        ("Pages requiring review", _count(results, lambda r: r.final_status == "REVIEW_REQUIRED")),
        ("Pages with errors", _count(results, lambda r: r.final_status == "ERROR")),
        ("Average quality score", round(mean(qualities), 2) if qualities else None),
        (
            "Pages below 150 DPI",
            _count(
                results,
                lambda r: (r.input_dpi is not None and r.input_dpi < 150)
                or (r.quality_warning or "").find("LOW_DPI") >= 0
                or "LOW_DPI" in (r.warning_list or []),
            ),
        ),
        (
            f"Pages below selected quality threshold ({quality_threshold:g})",
            _count(results, lambda r: r.quality_score is not None and r.quality_score < quality_threshold),
        ),
        ("Printed pages", _count(results, lambda r: r.document_type == "PRINTED")),
        ("Handwritten pages", _count(results, lambda r: r.document_type == "HANDWRITTEN")),
        ("Uncertain document type", _count(results, lambda r: r.document_type == "UNCERTAIN")),
        (
            "Rotated pages",
            _count(
                results,
                lambda r: r.rotation_angle is not None and abs(float(r.rotation_angle)) >= 0.2,
            ),
        ),
        ("Mirrored pages", _count(results, lambda r: r.mirror_corrected == "YES")),
        (
            "Deskewed pages",
            _count(results, lambda r: r.tilt_angle is not None and abs(float(r.tilt_angle)) >= 0.2),
        ),
        ("Rotation uncertain", _count(results, lambda r: r.rotation_status == "UNCERTAIN")),
        ("Mirror uncertain", _count(results, lambda r: r.mirror == "UNKNOWN")),
        ("Tilt uncertain", _count(results, lambda r: r.tilt_status == "UNCERTAIN")),
    ]

    header_font = Font(bold=True, color="FFFFFF", name="Calibri")
    ws.cell(3, 1, "Metric").fill = HEADER_FILL
    ws.cell(3, 1).font = header_font
    ws.cell(3, 2, "Value").fill = HEADER_FILL
    ws.cell(3, 2).font = header_font

    for i, (label, value) in enumerate(rows, start=4):
        a = ws.cell(i, 1, label)
        b = ws.cell(i, 2, value)
        a.border = THIN
        b.border = THIN
        if i % 2 == 0:
            a.fill = ALT_FILL
            b.fill = ALT_FILL
        if isinstance(value, float):
            b.number_format = "0.00"
        if "review" in label.lower() or "uncertain" in label.lower() or "below" in label.lower():
            if isinstance(value, (int, float)) and value:
                b.fill = REVIEW_FILL if "review" in label.lower() else WARN_FILL

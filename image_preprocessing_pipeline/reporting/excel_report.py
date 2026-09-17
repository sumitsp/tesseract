"""Professional Excel report: one row per page plus a Summary sheet."""

from __future__ import annotations

from pathlib import Path
from statistics import mean

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from image_preprocessing.results import PageResult

# Display columns only (pipeline fields unchanged).
COLUMNS: list[tuple[str, int]] = [
    ("Folder Name", 24),
    ("File Name", 32),
    ("Quality Score", 14),
    ("Quality", 10),
    ("Quality Warning", 28),
    ("Document Type", 22),
    ("Document Type Method", 22),
    ("Rotation Angle", 16),
    ("Tilt Angle", 14),
    ("Mirror", 12),
    ("Final Status", 18),
    ("Time taken", 14),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
ALT_FILL = PatternFill("solid", fgColor="F2F2F2")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
REVIEW_FILL = PatternFill("solid", fgColor="F8CBAD")
ERROR_FILL = PatternFill("solid", fgColor="F4C7C3")
OK_FILL = PatternFill("solid", fgColor="C6EFCE")
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
ANGLE_FORMAT = "0.00"
DPI_FORMAT = "0.0"
TIME_FORMAT = "0.000"


def _folder_name(result: PageResult) -> str:
    try:
        return Path(result.input_file).parent.name or ""
    except (TypeError, ValueError):
        return ""


def _file_name(result: PageResult) -> str:
    name = result.document_name or Path(result.input_file).name
    if int(result.page_number or 1) > 1:
        stem = Path(name).stem
        suffix = Path(name).suffix
        return f"{stem}_page_{int(result.page_number):03d}{suffix}"
    return name


def _quality_label(result: PageResult, quality_threshold: float) -> str:
    if result.quality_score is None:
        return "Bad"
    return "Good" if float(result.quality_score) >= quality_threshold else "Bad"


def _document_type_display(result: PageResult, quality_threshold: float) -> str:
    """Printed | Handwritten | Uncertain | Uncertain + Printed | Uncertain + Handwritten."""
    dt = (result.document_type or "").upper()
    method = (result.document_type_method or "").strip()
    p_hw = result.document_type_p_handwritten

    if dt == "PRINTED":
        return "Printed"
    if dt == "HANDWRITTEN":
        return "Handwritten"

    # Proper uncertain: blank, bad quality, classifier error — no lean label.
    if dt == "ERROR":
        return "Uncertain"
    if method == "blank_page":
        return "Uncertain"
    if result.quality_score is not None and float(result.quality_score) < quality_threshold:
        return "Uncertain"

    # Model could not commit: Uncertain + whichever side is closer (p_handwritten).
    if p_hw is not None:
        closer = "Handwritten" if float(p_hw) >= 0.5 else "Printed"
        return f"Uncertain + {closer}"
    return "Uncertain"


def _row_values(result: PageResult, quality_threshold: float) -> list:
    return [
        _folder_name(result),
        _file_name(result),
        result.quality_score,
        _quality_label(result, quality_threshold),
        result.quality_warning or "",
        _document_type_display(result, quality_threshold),
        result.document_type_method or "",
        result.rotation_angle,
        result.tilt_angle,
        result.mirror or "",
        result.final_status or "",
        result.processing_time_seconds,
    ]


def write_excel_report(path: Path, results: list[PageResult], quality_threshold: float) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    pages = wb.active
    pages.title = "Pages"
    _write_pages_sheet(pages, results, quality_threshold)
    summary = wb.create_sheet("Summary")
    _write_summary_sheet(summary, results, quality_threshold)
    wb.save(path)
    return path


def _write_pages_sheet(
    ws: Worksheet, results: list[PageResult], quality_threshold: float
) -> None:
    headers = [h for h, _ in COLUMNS]
    for col, (header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(1, col, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(2, len(results) + 1)}"
    ws.row_dimensions[1].height = 22

    angle_cols = {"Rotation Angle", "Tilt Angle"}
    numeric_cols = {"Quality Score"}

    for row_i, result in enumerate(results, start=2):
        values = _row_values(result, quality_threshold)
        for col, (header, _) in enumerate(COLUMNS, start=1):
            value = values[col - 1]
            cell = ws.cell(row_i, col, value)
            cell.border = THIN
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=header in {"Quality Warning", "File Name", "Folder Name"},
            )
            if row_i % 2 == 0:
                cell.fill = ALT_FILL
            if header in angle_cols and isinstance(value, (int, float)):
                cell.number_format = ANGLE_FORMAT
            elif header in numeric_cols and isinstance(value, (int, float)):
                cell.number_format = DPI_FORMAT
            elif header == "Time taken" and isinstance(value, (int, float)):
                cell.number_format = TIME_FORMAT

    last_row = max(2, len(results) + 1)
    last_col = get_column_letter(len(headers))
    status_col = get_column_letter(headers.index("Final Status") + 1)
    quality_col = get_column_letter(headers.index("Quality") + 1)
    doc_type_col = get_column_letter(headers.index("Document Type") + 1)
    mirror_col = get_column_letter(headers.index("Mirror") + 1)

    ws.conditional_formatting.add(
        f"{status_col}2:{status_col}{last_row}",
        CellIsRule(operator="equal", formula=['"REVIEW_REQUIRED"'], fill=REVIEW_FILL),
    )
    ws.conditional_formatting.add(
        f"{status_col}2:{status_col}{last_row}",
        CellIsRule(operator="equal", formula=['"ERROR"'], fill=ERROR_FILL),
    )
    ws.conditional_formatting.add(
        f"{status_col}2:{status_col}{last_row}",
        CellIsRule(operator="equal", formula=['"CORRECTED"'], fill=OK_FILL),
    )
    ws.conditional_formatting.add(
        f"{quality_col}2:{quality_col}{last_row}",
        CellIsRule(operator="equal", formula=['"Bad"'], fill=WARN_FILL),
    )
    ws.conditional_formatting.add(
        f"{doc_type_col}2:{doc_type_col}{last_row}",
        FormulaRule(formula=[f'ISNUMBER(SEARCH("Uncertain",{doc_type_col}2))'], fill=WARN_FILL),
    )
    ws.conditional_formatting.add(
        f"{mirror_col}2:{mirror_col}{last_row}",
        CellIsRule(operator="equal", formula=['"UNKNOWN"'], fill=WARN_FILL),
    )
    warn_col = get_column_letter(headers.index("Quality Warning") + 1)
    ws.conditional_formatting.add(
        f"A2:{last_col}{last_row}",
        FormulaRule(formula=[f'LEN(${warn_col}2)>0'], fill=WARN_FILL),
    )


def _count(results: list[PageResult], pred) -> int:
    return sum(1 for r in results if pred(r))


def _write_summary_sheet(ws: Worksheet, results: list[PageResult], quality_threshold: float) -> None:
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18
    title = ws.cell(1, 1, "Preprocessing summary")
    title.font = Font(bold=True, size=14, color="1F4E79", name="Calibri")
    ws.merge_cells("A1:B1")

    docs = {r.document_name for r in results}
    qualities = [r.quality_score for r in results if r.quality_score is not None]
    rows = [
        ("Total documents", len(docs)),
        ("Total pages", len(results)),
        (
            "Pages successfully processed",
            _count(results, lambda r: r.final_status in {"CORRECTED", "UNCHANGED"}),
        ),
        ("Pages requiring review", _count(results, lambda r: r.final_status == "REVIEW_REQUIRED")),
        ("Pages with errors", _count(results, lambda r: r.final_status == "ERROR")),
        ("Average quality score", round(mean(qualities), 2) if qualities else None),
        (
            "Pages below 150 DPI",
            _count(
                results,
                lambda r: (r.input_dpi is not None and r.input_dpi < 150)
                or (r.quality_warning or "").find("LOW_DPI") >= 0
                or "LOW_DPI" in (r.warning_list or []),
            ),
        ),
        (
            f"Pages below selected quality threshold ({quality_threshold:g})",
            _count(results, lambda r: r.quality_score is not None and r.quality_score < quality_threshold),
        ),
        ("Printed pages", _count(results, lambda r: r.document_type == "PRINTED")),
        ("Handwritten pages", _count(results, lambda r: r.document_type == "HANDWRITTEN")),
        ("Uncertain document type", _count(results, lambda r: r.document_type == "UNCERTAIN")),
        (
            "Rotated pages",
            _count(
                results,
                lambda r: r.rotation_angle is not None and abs(float(r.rotation_angle)) >= 0.2,
            ),
        ),
        ("Pages flagged mirrored (not flipped)", _count(results, lambda r: r.mirror == "YES")),
        (
            "Deskewed pages",
            _count(results, lambda r: r.tilt_status == "APPLIED"),
        ),
        ("Rotation uncertain", _count(results, lambda r: r.rotation_status == "UNCERTAIN")),
        ("Rotation rejected by validation", _count(results, lambda r: r.rotation_status == "REJECTED")),
        ("Mirror uncertain", _count(results, lambda r: r.mirror == "UNKNOWN")),
        ("Tilt uncertain", _count(results, lambda r: r.tilt_status == "UNCERTAIN")),
        (
            "Tilt withheld (orientation unconfirmed)",
            _count(results, lambda r: r.tilt_status == "NOT_APPLIED"),
        ),
    ]

    header_font = Font(bold=True, color="FFFFFF", name="Calibri")
    ws.cell(3, 1, "Metric").fill = HEADER_FILL
    ws.cell(3, 1).font = header_font
    ws.cell(3, 2, "Value").fill = HEADER_FILL
    ws.cell(3, 2).font = header_font

    for i, (label, value) in enumerate(rows, start=4):
        a = ws.cell(i, 1, label)
        b = ws.cell(i, 2, value)
        a.border = THIN
        b.border = THIN
        if i % 2 == 0:
            a.fill = ALT_FILL
            b.fill = ALT_FILL
        if isinstance(value, float):
            b.number_format = "0.00"
        if "review" in label.lower() or "uncertain" in label.lower() or "below" in label.lower():
            if isinstance(value, (int, float)) and value:
                b.fill = REVIEW_FILL if "review" in label.lower() else WARN_FILL
